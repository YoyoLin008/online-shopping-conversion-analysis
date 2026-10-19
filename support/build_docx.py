"""Create an editable Word version of the checked academic PDF."""
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = Inches(0.75)
section.top_margin = Pt(42)
section.bottom_margin = Pt(54)
section.footer_distance = Pt(22)

def font(style, family, size, bold=False):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + key), family)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + key), None)

font(doc.styles["Normal"], "Times New Roman", 11)
normal = doc.styles["Normal"].paragraph_format
normal.space_before = normal.space_after = Pt(0)
normal.line_spacing = Pt(13)
normal.widow_control = False

font(doc.styles["Title"], "Times New Roman", 14, True)
title_style = doc.styles["Title"].paragraph_format
title_style.space_before = Pt(0)
title_style.space_after = Pt(2)
title_style.line_spacing = Pt(17)
title_style.keep_with_next = True
for element in doc.styles["Title"].element.xpath("./w:pPr/w:pBdr | ./w:rPr/w:spacing | ./w:rPr/w:kern"):
    element.getparent().remove(element)

for name, family, size, leading in [
    ("Code", "Courier New", 9.5, 11),
    ("Console", "Courier New", 9.5, 11.5),
    ("Figure Note", "Times New Roman", 10, 12),
]:
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    font(style, family, size)
    fmt = style.paragraph_format
    fmt.line_spacing = Pt(leading)
    fmt.keep_together = False
    fmt.keep_with_next = False
    fmt.widow_control = False

title = doc.add_paragraph("Purchase Conversion by Visitor Type", "Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
name = doc.add_paragraph("Yunya Lin")
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.line_spacing = Pt(14)
name.paragraph_format.space_after = Pt(12)
intro = doc.add_paragraph(
    "Does month composition explain the purchase-rate difference between new "
    "and returning visitors? The UCI data contain 12,330 sessions from distinct "
    "users. Revenue records whether a purchase occurred, not its monetary value."
)
intro.paragraph_format.space_after = Pt(6)

lines = (ROOT / "coding_essay.py").read_text().splitlines()
split = next(i for i, line in enumerate(lines)
             if line.startswith("# Intervals describe"))
doc.add_paragraph("\n".join(lines[:split]).rstrip("\n"), "Code")
p = doc.add_paragraph("\n".join(lines[split:]), "Code")
p.paragraph_format.page_break_before = True
p.paragraph_format.space_after = Pt(8)

p = doc.add_paragraph()
p.add_run("Output").bold = True
p.paragraph_format.keep_with_next = True
p.paragraph_format.space_after = Pt(3)
stdout = (ROOT / "support/results/console_output.txt").read_text()
p = doc.add_paragraph(stdout.rstrip("\n"), "Console")
p.paragraph_format.space_after = Pt(9)

p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.0
p.paragraph_format.keep_with_next = True
p.paragraph_format.space_after = Pt(6)
picture = p.add_run().add_picture(
    str(ROOT / "support/results/conversion_comparison.png"), width=Inches(7)
)
picture._inline.docPr.set("descr",
    "Observed purchase rates for New and Returning visitors with 95 percent "
    "Wilson intervals, alongside open diamonds for month-standardized rates.")
caption = doc.add_paragraph(
    "Figure 1. Purchase rates by visitor type. Circles and whiskers show "
    "observed rates and 95% Wilson intervals; open diamonds show point estimates "
    "using the pooled monthly distribution of the retained sample.", "Figure Note"
)
caption.paragraph_format.space_after = Pt(10)
interpretation = doc.add_paragraph(
    "The observed gap is 10.98 percentage points; giving both groups the same "
    "month mix leaves a 10.67-point gap. Month composition alone therefore "
    "does not account for the difference in this sample. This comparison is "
    "descriptive, not causal: visitor characteristics may differ, and February "
    "has only one new visitor. No uncertainty interval is estimated for the "
    "month-standardized rates."
)
interpretation.paragraph_format.space_after = Pt(10)
doc.add_paragraph(
    "Source: Sakar, C., and Kastro, Y. (2018). Online Shoppers Purchasing "
    "Intention Dataset. UCI Machine Learning Repository. "
    "https://doi.org/10.24432/C5F88Q. CC BY 4.0.", "Figure Note"
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.line_spacing = Pt(12)
run = footer.add_run()
run.font.name = "Times New Roman"
run.font.size = Pt(10)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
field_run = OxmlElement("w:r")
value = OxmlElement("w:t")
value.text = "1"
field_run.append(value)
field.append(field_run)
run._r.addnext(field)

doc.core_properties.title = "Purchase Conversion by Visitor Type"
doc.core_properties.author = "Yunya Lin"
doc.core_properties.subject = "Python analysis of UCI online shopping sessions"
doc.core_properties.comments = ""
output = ROOT / "Yunya_Lin_UChicago_Coding_Essay.docx"
doc.save(output)
print(output)
